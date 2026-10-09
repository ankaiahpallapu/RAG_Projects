from __future__ import annotations

import logging
import os
import re
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path
from typing import Iterator

from .config import Config

log = logging.getLogger(__name__)

SUPPORTED = {".txt", ".md", ".rst", ".csv", ".pdf", ".docx", ".html", ".htm"}


# --------------------------------------------------------------------------- loading
def read_file(path: Path) -> str:
    """Extract plain text from a file. Returns '' (and logs) if the file can't be parsed."""
    ext = path.suffix.lower()
    try:
        if ext in {".txt", ".md", ".rst", ".csv"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            pages = []
            for page in reader.pages:
                try:
                    pages.append(page.extract_text() or "")
                except Exception:  # one bad page shouldn't lose the whole document
                    continue
            return "\n\n".join(pages)
        if ext == ".docx":
            from docx import Document

            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n\n".join(parts)
        if ext in {".html", ".htm"}:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
            for tag in soup(["script", "style", "noscript"]):
                tag.decompose()
            return soup.get_text("\n")
    except Exception as exc:  # noqa: BLE001 - ingestion must be resilient to corrupt files
        log.warning("failed to read %s: %s", path, exc)
    return ""


# --------------------------------------------------------------------------- chunking
_SENT = re.compile(r"(?<=[.!?])\s+")


def _overlap_tail(text: str, overlap: int) -> str:
    if overlap <= 0 or len(text) <= overlap:
        return ""
    tail = text[-overlap:]
    sp = tail.find(" ")  # start the overlap on a word boundary
    return tail[sp + 1 :] if sp != -1 else tail


def _pack(pieces: list[str], size: int, overlap: int, sep: str) -> list[str]:
    out: list[str] = []
    buf = ""
    for piece in pieces:
        if not buf:
            buf = piece
        elif len(buf) + len(sep) + len(piece) <= size:
            buf = buf + sep + piece
        else:
            out.append(buf)
            tail = _overlap_tail(buf, overlap)
            buf = f"{tail}{sep}{piece}" if tail else piece
    if buf:
        out.append(buf)
    return out


def _sentences(paragraph: str, size: int) -> list[str]:
    pieces: list[str] = []
    for sent in _SENT.split(paragraph):
        if len(sent) <= size:
            pieces.append(sent)
        else:  # pathological sentence with no punctuation: hard cut
            pieces.extend(sent[i : i + size] for i in range(0, len(sent), size))
    return pieces


def split_text(text: str, size: int, overlap: int) -> list[str]:
    """Paragraph-aware greedy chunking with sentence fallback and word-aligned overlap."""
    text = text.replace("\r\n", "\n")
    paragraphs = [re.sub(r"[ \t]+", " ", p).strip() for p in re.split(r"\n\s*\n", text)]
    paragraphs = [p for p in paragraphs if p]

    out: list[str] = []
    buf = ""
    for para in paragraphs:
        if len(para) > size:
            if buf:
                out.append(buf)
                buf = ""
            out.extend(_pack(_sentences(para, size), size, overlap, " "))
        elif not buf:
            buf = para
        elif len(buf) + 2 + len(para) <= size:
            buf += "\n\n" + para
        else:
            tail = _overlap_tail(buf, overlap)
            out.append(buf)
            buf = f"{tail}\n\n{para}" if tail else para
    if buf:
        out.append(buf)
    return out


# --------------------------------------------------------------------------- directory ingestion
def _process_file(args: tuple[str, str, int, int, int]) -> list[tuple[str, int, str]]:
    path, root, size, overlap, min_chars = args
    p = Path(path)
    text = read_file(p)
    if not text.strip():
        return []
    doc_id = p.relative_to(root).as_posix()
    chunks = [c for c in split_text(text, size, overlap) if len(c) >= min_chars]
    return [(doc_id, i, c) for i, c in enumerate(chunks)]


def iter_files(root: str | Path) -> Iterator[Path]:
    for dirpath, _, files in os.walk(root):
        for name in sorted(files):
            if Path(name).suffix.lower() in SUPPORTED:
                yield Path(dirpath) / name


def iter_chunks(root: str | Path, cfg: Config) -> Iterator[tuple[str, int, str]]:
    """Yield (doc_id, position, text) for every chunk under `root`, parsing files in parallel."""
    root = str(Path(root).resolve())
    tasks = (
        (str(p), root, cfg.chunk_size, cfg.chunk_overlap, cfg.min_chunk_chars) for p in iter_files(root)
    )
    if cfg.ingest_workers <= 1:
        for task in tasks:
            yield from _process_file(task)
        return
    with ProcessPoolExecutor(max_workers=cfg.ingest_workers) as pool:
        for result in pool.map(_process_file, tasks, chunksize=16):
            yield from result
